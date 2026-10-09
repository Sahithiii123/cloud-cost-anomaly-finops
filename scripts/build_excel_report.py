"""Build an Excel report with formula summaries and native Excel charts."""
import json
import os
from pathlib import Path
import subprocess
import shutil

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parent.parent


def append_frame(sheet, frame):
    sheet.append([str(column) for column in frame.columns])
    for row in frame.itertuples(index=False, name=None):
        sheet.append([value.to_pydatetime() if hasattr(value, "to_pydatetime") else value for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def style_sheet(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244C5A")
    for column_cells in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12), 36)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def main():
    report_path = ROOT / "reports" / "monthly_finops_report.xlsx"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    raw = pd.read_csv(ROOT / "data" / "cur_sample.csv", parse_dates=["usage_date"])
    anomalies = pd.read_csv(ROOT / "outputs" / "synthetic" / "anomalies_detected.csv")
    tags = pd.read_csv(ROOT / "outputs" / "synthetic" / "tag_compliance_by_month.csv")
    tag_detail = pd.read_csv(ROOT / "outputs" / "synthetic" / "tag_compliance_by_service_region_month.csv")
    forecast_daily = pd.read_csv(ROOT / "outputs" / "synthetic" / "forecast_daily.csv", parse_dates=["usage_date"])
    summary = json.loads((ROOT / "outputs" / "synthetic" / "analysis_summary.json").read_text(encoding="utf-8"))

    workbook = Workbook()
    raw_sheet = workbook.active
    raw_sheet.title = "Raw Data"
    append_frame(raw_sheet, raw)
    raw_table = Table(displayName="RawCURData", ref=raw_sheet.dimensions)
    raw_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    raw_sheet.add_table(raw_table)

    months = sorted(pd.to_datetime(raw["usage_date"]).dt.to_period("M").astype(str).unique())
    monthly = workbook.create_sheet("Monthly Summary")
    monthly.append(["Usage month", "Actual spend (USD)"])
    for row_num, month in enumerate(months, 2):
        monthly.append([month, f'=SUMIFS(\'Raw Data\'!$E:$E,\'Raw Data\'!$A:$A,">="&DATEVALUE(A{row_num}&"-01"),\'Raw Data\'!$A:$A,"<"&EDATE(DATEVALUE(A{row_num}&"-01"),1))'])

    service_summary = workbook.create_sheet("Service Summary")
    service_summary.append(["Service", "Actual spend (USD)"])
    for row_num, service in enumerate(sorted(raw["service"].unique()), 2):
        service_summary.append([service, f"=SUMIF('Raw Data'!$B:$B,A{row_num},'Raw Data'!$E:$E)"])

    region_summary = workbook.create_sheet("Region Summary")
    region_summary.append(["Region", "Actual spend (USD)"])
    for row_num, region in enumerate(sorted(raw["region"].unique()), 2):
        region_summary.append([region, f"=SUMIF('Raw Data'!$C:$C,A{row_num},'Raw Data'!$E:$E)"])

    budget = workbook.create_sheet("Budget vs Actual")
    budget.append(["Usage month", "Actual spend (USD)", "Budget input (USD)", "Variance (USD)"])
    for row_num, month in enumerate(months, 2):
        budget.append([month, f"='Monthly Summary'!B{row_num}", None, f'=IF(C{row_num}="","",B{row_num}-C{row_num})'])
    budget[1][2].comment = __import__("openpyxl").comments.Comment("Enter an approved budget; no budget values are invented by this report.", "FinOps demo")

    anomaly_sheet = workbook.create_sheet("Anomalies")
    append_frame(anomaly_sheet, anomalies)
    tag_sheet = workbook.create_sheet("Tag Compliance")
    append_frame(tag_sheet, tags)
    tag_detail_sheet = workbook.create_sheet("Tag Detail")
    append_frame(tag_detail_sheet, tag_detail)
    forecast = workbook.create_sheet("Forecast")
    forecast.append(["Metric", "Value", "Unit / note"])
    forecast.append(["Regression next 30 days", summary["regression_forecast_30d_usd"], "USD; forecast, not actual spend"])
    forecast.append(["Exponential smoothing next 30 days", summary["exponential_smoothing_forecast_30d_usd"], "USD; forecast, not actual spend"])
    forecast.append(["Method agreement", summary["forecast_agreement_pct"], "Percent difference of forecast totals"])
    forecast.append(["Regression rolling-origin MAPE", summary["forecast_regression_mape_pct"], "Percent; mean of 30-day-origin backtests"])
    forecast.append(["Exponential smoothing rolling-origin MAPE", summary["forecast_exponential_smoothing_mape_pct"], "Percent; mean of 30-day-origin backtests"])
    forecast_daily_sheet = workbook.create_sheet("Forecast Daily")
    append_frame(forecast_daily_sheet, forecast_daily)

    spend_chart = LineChart()
    spend_chart.title = "Monthly Actual Spend"
    spend_chart.y_axis.title = "USD"
    spend_chart.x_axis.title = "Month"
    spend_chart.add_data(Reference(monthly, min_col=2, min_row=1, max_row=monthly.max_row), titles_from_data=True)
    spend_chart.set_categories(Reference(monthly, min_col=1, min_row=2, max_row=monthly.max_row))
    monthly.add_chart(spend_chart, "D2")
    tag_chart = BarChart()
    tag_chart.title = "Monthly Tag Compliance"
    tag_chart.y_axis.title = "Percent"
    tag_chart.add_data(Reference(tag_sheet, min_col=7, min_row=1, max_row=tag_sheet.max_row), titles_from_data=True)
    tag_chart.set_categories(Reference(tag_sheet, min_col=1, min_row=2, max_row=tag_sheet.max_row))
    tag_sheet.add_chart(tag_chart, "J2")

    for sheet in workbook.worksheets:
        style_sheet(sheet)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(report_path)
    print(f"Created {report_path.relative_to(ROOT)} (formula summaries, not native PivotTables).")
    if os.name == "nt":
        powershell = shutil.which("powershell.exe") or shutil.which("powershell")
        com_script = ROOT / "scripts" / "build_excel_report_com.ps1"
        if powershell and com_script.is_file():
            result = subprocess.run([powershell, "-NoProfile", "-ExecutionPolicy", "Bypass",
                                     "-File", str(com_script)], cwd=ROOT, check=False)
            if result.returncode:
                print("Excel COM XLSM build failed; the formula-driven XLSX remains available.")
        else:
            print("Excel COM is unavailable; the formula-driven XLSX remains available.")


if __name__ == "__main__":
    main()